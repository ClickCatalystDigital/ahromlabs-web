#!/usr/bin/env python3
"""Writes public/boq/interior-boq-template.xlsx and .csv from the item library
(src/lib/boq.ts), so the template can't drift from /boq/items. Run after
changing the library:  python3 scripts/build-boq-template.py
Needs openpyxl."""
import csv, json, subprocess
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

root = Path(__file__).resolve().parent.parent
data = json.loads(subprocess.run(["npx", "jiti", "scripts/boq-json.ts"], cwd=root, capture_output=True, text=True, check=True).stdout)
out = root / "public" / "boq"
out.mkdir(parents=True, exist_ok=True)

HEAD = ["S.No", "Description", "Unit", "Measurement basis", "Qty", "Rate", "Amount", "Item id"]
rows = []  # (kind, values)
n = 0
for t in data["trades"]:
    rows.append(("trade", [None, t["name"].upper()]))
    for i in (i for i in data["items"] if i["trade"] == t["slug"]):
        n += 1
        rows.append(("item", [n, i["description"], i["unitLabel"], i["measure"], None, None, None, i["id"]]))

wb = Workbook()
ws = wb.active
ws.title = "BOQ"
ws.append(["Interior BOQ"])
ws.append(["Fill the [brackets], quantities and rates. From the item library at ahromlabs.com/boq/items (free under CC BY 4.0, credit Ahrom Labs); check a filled BOQ at ahromlabs.com/boq/check."])
ws.append([])
ws.append(HEAD)
head_row = ws.max_row
first = head_row + 1
for kind, values in rows:
    ws.append(values)
    r = ws.max_row
    if kind == "trade":
        ws.cell(r, 2).font = Font(bold=True)
    else:
        ws.cell(r, 7).value = f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),E{r}*F{r},"")'
        ws.cell(r, 7).number_format = "#,##,##0.00"
        ws.cell(r, 6).number_format = "#,##,##0.00"
last = ws.max_row
ws.append([None, "Total (before GST)", None, None, None, None, f"=SUM(G{first}:G{last})"])
ws.cell(ws.max_row, 2).font = Font(bold=True)
ws.cell(ws.max_row, 7).font = Font(bold=True)
ws.cell(ws.max_row, 7).number_format = "#,##,##0.00"

ws["A1"].font = Font(bold=True, size=14)
for c in ws[head_row]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="EEEEEE")
for col, width in zip("ABCDEFGH", [6, 70, 11, 45, 8, 11, 13, 26]):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=first, max_row=last):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = ws.cell(first, 1)
wb.save(out / "interior-boq-template.xlsx")

with open(out / "interior-boq-template.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(HEAD)
    for kind, values in rows:
        w.writerow((values + [None] * len(HEAD))[: len(HEAD)])

print(f"boq template: {n} items -> {out.relative_to(root)}/interior-boq-template.(xlsx|csv)")
