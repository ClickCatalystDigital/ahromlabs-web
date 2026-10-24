#!/usr/bin/env python3
"""Writes public/erp/erp-requirements-template.xlsx and .csv from the requirements library
(src/lib/erp.ts), so the template can't drift from /erp/requirements. Run after changing the
library:  python3 scripts/build-erp-template.py
Needs openpyxl.

The sheet is an RFP and a scorecard in one: priorities and vendor answers are dropdowns, and each
line scores priority weight (Must 3, Should 2, Could 1) x answer (Yes 1, Partly 0.5, No 0)."""
import csv, json, subprocess
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

root = Path(__file__).resolve().parent.parent
data = json.loads(subprocess.run(["npx", "jiti", "scripts/erp-json.ts"], cwd=root, capture_output=True, text=True, check=True).stdout)
out = root / "public" / "erp"
out.mkdir(parents=True, exist_ok=True)

HEAD = ["S.No", "Department", "Requirement", "Priority", "What the vendor must show", "Law", "Vendor answer", "Vendor evidence", "Score", "Requirement id"]
reqs = data["requirements"]

wb = Workbook()
ws = wb.active
ws.title = "Requirements"
ws.append(["ERP requirements"])
ws.append([f"From the ERP requirements checklist v{data['version']} at ahromlabs.com/erp/requirements (free under CC BY 4.0, credit Ahrom Labs). Set each priority, send to vendors, then fill their answers to score them."])
ws.append([])
ws.append(HEAD)
head_row = ws.max_row
first = head_row + 1
dept = None
for n, r in enumerate(reqs, 1):
    if r["department"] != dept:
        dept = r["department"]
    law = r.get("governedBy", {}).get("law", "")
    ws.append([n, dept, r["requirement"], r["priority"], r["test"], law, None, None, None, r["id"]])
    row = ws.max_row
    ws.cell(row, 9).value = f'=IF(G{row}="","",IF(D{row}="Must",3,IF(D{row}="Should",2,1))*IF(G{row}="Yes",1,IF(G{row}="Partly",0.5,0)))'
last = ws.max_row
ws.append([None, "Score", None, None, None, None, None, "Out of", f"=SUM(I{first}:I{last})", None])
tot = ws.max_row
ws.cell(tot, 8).value = f'=SUMPRODUCT((G{first}:G{last}<>"")*((D{first}:D{last}="Must")*3+(D{first}:D{last}="Should")*2+(D{first}:D{last}="Could")*1))'
ws.append([None, "Score as % of the lines answered", None, None, None, None, None, None, f'=IF(H{tot}=0,"",I{tot}/H{tot})', None])
pct = ws.max_row
ws.cell(pct, 9).number_format = "0%"
for r_ in (tot, pct):
    ws.cell(r_, 2).font = Font(bold=True)
    ws.cell(r_, 9).font = Font(bold=True)

prio = DataValidation(type="list", formula1='"Must,Should,Could"', allow_blank=False)
ans = DataValidation(type="list", formula1='"Yes,Partly,No"', allow_blank=True)
ws.add_data_validation(prio)
ws.add_data_validation(ans)
prio.add(f"D{first}:D{last}")
ans.add(f"G{first}:G{last}")

ws["A1"].font = Font(bold=True, size=14)
for c in ws[head_row]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="EEEEEE")
for col, width in zip("ABCDEFGHIJ", [6, 22, 60, 10, 55, 30, 14, 40, 8, 30]):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=first, max_row=last):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = ws.cell(first, 1)

how = wb.create_sheet("How to use")
for line in [
    "1. Delete the lines that don't apply, or build a tailored list at ahromlabs.com/erp/check.",
    "2. Set each line's Priority: Must, Should or Could.",
    "3. Send the sheet to each vendor. Ask them to answer Yes, Partly or No, and to show the test in a demo.",
    "4. Fill each vendor's answers in a copy of this sheet. The Score column and the total compare them.",
    "5. A 'Yes' you haven't seen demonstrated counts as 'Partly'.",
    "",
    "Lines with a Law are required of the business by the provision named. The quoted text and source for each are at ahromlabs.com/erp/requirements.",
]:
    how.append([line])
how.column_dimensions["A"].width = 110
wb.save(out / "erp-requirements-template.xlsx")

with open(out / "erp-requirements-template.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow([h for h in HEAD if h != "Score"])
    for n, r in enumerate(reqs, 1):
        w.writerow([n, r["department"], r["requirement"], r["priority"], r["test"], r.get("governedBy", {}).get("law", ""), "", "", r["id"]])

print(f"erp template: {len(reqs)} requirements -> {out.relative_to(root)}/erp-requirements-template.(xlsx|csv)")
